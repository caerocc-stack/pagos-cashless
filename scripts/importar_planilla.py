"""
Importa al sistema la planilla de Google Sheets "ALUMNOS Y HORARIOS" (exportada a .xlsx):
cursos y valores, alumnos (incluidos los que cursan más de un curso), pagos del año, deuda del año
anterior, matrícula ya paga y el control horario del personal.

    python scripts/importar_planilla.py planilla.xlsx [--reemplazar] [--sin-horario]

--reemplazar borra antes los alumnos, pagos, cursos por alumno, saldos y horario existentes.
"""
import hashlib
import os
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, time
from decimal import Decimal

import openpyxl
from dotenv import load_dotenv

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
load_dotenv(os.path.join(RAIZ, ".env"))

from app.database import SessionLocal  # noqa: E402
import app.esquema  # noqa: E402,F401
from app.models.alumno import Alumno  # noqa: E402
from app.models.curso import Curso, AlumnoCurso  # noqa: E402
from app.models.pago_cuota import PagoCuota  # noqa: E402
from app.models.saldo import Saldo  # noqa: E402
from app.models.horario import Empleado, RegistroHorario, LiquidacionHoras  # noqa: E402
from app.routers.cursos import curso_de_alumno  # noqa: E402

ANIO = 2026
MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]

# Cursos según PARAMETROS (nombre, área, patrón, cuota, n_cuotas, matrícula, mes de inicio o None = inferir)
CURSOS = [
    ("Ciclo Básico", "Ciclo Basico", None, 45000, 10, 50000, 3, True),
    ("Mecánica", "Mecanica", None, 45000, 10, 50000, 3, True),
    ("Aviónica", "Avionica", None, 45000, 10, 50000, 3, True),
    ("MMH (Aeronavegantes)", "Aeronavegantes", "MMH", 45000, 10, 50000, 3, False),
    ("ORR A (Examen)", "Examenes", "ORR A", 180000, 3, 0, None, False),
    ("ORR B (Examen)", "Examenes", "ORR B", 180000, 3, 0, 8, False),   # curso de agosto/septiembre
    ("RAAC 147 (Examen)", "Examenes", "RAAC", 100000, 1, 0, 6, False),  # examen del 16/06
    ("DESP AER", "Aeronavegantes", "DESP AER", 50000, 1, 0, None, False),
    ("TCP", "Aeronavegantes", "TCP", 50000, 1, 0, None, False),
    ("PCA-HVI", "Aeronavegantes", "PCA-HVI", 50000, 1, 0, None, False),
]
TIPO_DIA = {"LABORAL": "Normal", "VACACIONES": "Vacaciones", "FERIADO": "Feriado", "LICENCIA": "Licencia",
            "FRANCO": "Franco", "AUSENTE": "Ausente", "ENFERMEDAD": "Licencia", "NO LABORABLE": "Franco"}


def txt(v) -> str:
    return "" if v is None else str(v).strip()


def num_str(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\D", "", str(v))


def limpiar_emails(v) -> str | None:
    """Corrige errores de tipeo comunes (gmail,com · espacios · mayúsculas) y deja los correos separados por '; '.
    Con ese correo y el DNI del alumno la familia entra al portal."""
    t = txt(v).lower().replace("mailto:", "")
    t = re.sub(r"@([a-z0-9-]+),(com|ar|net|org)\b", r"@\1.\2", t)   # gmail,com -> gmail.com
    t = re.sub(r"\.con\b", ".com", t)                                 # hotmail.con -> hotmail.com
    if "@" not in t:  # delmastro42gmail.com -> delmastro42@gmail.com
        t = re.sub(r"^([a-z0-9._%+-]+?)(gmail|hotmail|yahoo|outlook|live|icloud)\.com", r"\1@\2.com", t.strip())
    correos = [c.strip(" .,;") for c in re.split(r"[;\s/]+|,(?=\s*[a-z0-9._%+-]+@)", t) if "@" in c]
    validos = [c for c in correos if re.fullmatch(r"[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}", c)]
    return "; ".join(dict.fromkeys(validos))[:150] or None


def separar_nombre(completo: str) -> tuple[str, str]:
    """'DI NÁPOLI Agustina Celeste' -> ('DI NÁPOLI', 'Agustina Celeste'). Admite 'APELLIDO, Nombre'."""
    s = re.sub(r"\s+becad[oa]\s*$", "", completo.strip(), flags=re.I)
    s = re.sub(r"\s+", " ", s)
    if "," in s:
        ape, nom = s.split(",", 1)
        return ape.strip(), nom.strip()
    partes = s.split(" ")
    ape = []
    for p in partes:
        if p.upper() == p and any(ch.isalpha() for ch in p):
            ape.append(p)
        else:
            break
    if not ape or len(ape) == len(partes):  # todo en mayúsculas o sin mayúsculas: primera palabra = apellido
        ape = partes[:1] if len(partes) > 1 else partes
    return " ".join(ape), " ".join(partes[len(ape):]) or "-"


def division_de(curso: str) -> str | None:
    m = re.search(r"°\s*([A-Z])\b", curso) or re.search(r"\s([AB])$", curso)
    return m.group(1) if m else None


def a_fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def a_hora(v) -> time | None:
    """900 -> 09:00 ; 2030 -> 20:30 ; también acepta time/datetime."""
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        return None
    h, m = divmod(n, 100)
    return time(h, m) if 0 <= h < 24 and 0 <= m < 60 else None


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    db = SessionLocal()
    reemplazar = "--reemplazar" in sys.argv

    if db.query(Alumno).count() and not reemplazar:
        sys.exit("Ya hay alumnos cargados. Usá --reemplazar para borrarlos y volver a importar.")
    if reemplazar:
        from app.models.portal import CuponEnviado, Consulta
        db.query(AlumnoCurso).delete()
        db.query(PagoCuota).delete()
        db.query(CuponEnviado).delete()
        db.query(Consulta).update({"alumno_id": None})
        db.query(Saldo).delete()
        db.query(Alumno).delete()
        if "--sin-horario" not in sys.argv:
            db.query(RegistroHorario).delete()
            db.query(LiquidacionHoras).delete()
            db.query(Empleado).delete()
        db.flush()

    # ------------------------------------------------------------ filas de ALUMNOS 2026
    filas = [f for f in list(wb["ALUMNOS 2026"].iter_rows(values_only=True))[3:]
             if f and txt(f[1]) and txt(f[5]) and num_str(f[2]) not in ("", "99999999")]

    # Mes de inicio de los cursos de pago único / ORR A: el primer mes en que se registran pagos
    primer_mes = defaultdict(lambda: 12)
    for f in filas:
        for m in range(12):
            if f[8 + 2 * m]:
                for nombre, _, patron, *_ in CURSOS:
                    if patron and txt(f[5]).upper().startswith(patron.upper()):
                        primer_mes[patron] = min(primer_mes[patron], m + 1)

    # ------------------------------------------------------------ cursos
    existentes = db.query(Curso).all()
    cursos = {}
    for nombre, area, patron, cuota, n, mat, mes, viejas in CURSOS:
        c = next((x for x in existentes if (patron and x.patron == patron) or (not patron and not x.patron and x.area == area)), None)
        if not c and patron == "ORR A":
            c = next((x for x in existentes if x.patron == "ORR"), None)  # el "ORR" viejo pasa a ser ORR A
        if not c:
            c = Curso(nombre=nombre)
            db.add(c)
        c.nombre, c.area, c.patron, c.cuota, c.n_cuotas, c.matricula = nombre, area, patron, Decimal(cuota), n, Decimal(mat)
        inferido = primer_mes.get(patron) if patron else None
        c.mes_inicio = mes or (max(3, inferido) if inferido else 3)
        c.cuota_anio_anterior = Decimal(30000) if viejas else Decimal(0)
        c.cobra_cuotas_viejas, c.anio, c.activo = viejas, ANIO, True
        cursos[nombre] = c
    for x in existentes:  # cursos que ya no están en la planilla
        if x not in cursos.values():
            x.activo = False
    db.flush()
    lista_cursos = list(cursos.values())
    print("Cursos:", ", ".join(f"{c.nombre} (${c.cuota:,.0f}×{c.n_cuotas}{' +mat' if c.matricula else ''}, desde mes {c.mes_inicio})" for c in lista_cursos))

    # ------------------------------------------------------------ alumnos (agrupados por DNI)
    grupos: "OrderedDict[str, list]" = OrderedDict()
    for f in filas:
        grupos.setdefault(num_str(f[2]), []).append(f)
    legajos_usados = set()
    alumnos = {}
    sin_curso = []
    for dni, fs in grupos.items():
        f = fs[0]
        apellido, nombre = separar_nombre(txt(f[1]))
        legajo = num_str(f[3])
        if not legajo or legajo in legajos_usados:
            legajo = "D" + dni
        legajos_usados.add(legajo)
        cond = txt(f[31]) or "Regular"
        deuda = max((Decimal(str(x[43])) for x in fs if isinstance(x[43], (int, float)) and x[43] > 0), default=None)
        a = Alumno(legajo=legajo, dni=dni[:10], apellido=apellido[:100], nombre=nombre[:100], curso=txt(f[5])[:20],
                   area=txt(f[4]) or None, division=division_de(txt(f[5])), email=limpiar_emails(f[68]),
                   telefono=num_str(f[67])[:40] or None, fecha_nacimiento=a_fecha(f[66]),
                   condicion=cond if cond in ("Regular", "Becado", "Exento", "Baja") else "Regular",
                   modalidad=txt(f[32]) or "Mensual", deuda_anio_anterior=deuda,
                   matricula_paga=txt(f[42]).lower() == "matriculado"           # AQ: matrícula manual
                   or (txt(f[53]).lower() == "no" and not f[46]),                # BB "Debe matrícula" = No sin pago 2026 (AU)
                   activo=cond != "Baja")
        db.add(a)
        alumnos[dni] = (a, fs)
    db.flush()
    for a, fs in alumnos.values():
        db.add(Saldo(alumno_id=a.id, monto=Decimal("0")))

    # Cursos de cada alumno: los que cursan más de uno quedan con sus cursos explícitos
    multi = 0
    curso_de_fila = {}
    for dni, (a, fs) in alumnos.items():
        for f in fs:
            c = curso_de_alumno(Alumno(curso=txt(f[5]), area=txt(f[4])), lista_cursos)
            curso_de_fila[id(f)] = c
            if not c:
                sin_curso.append((a.apellido, txt(f[5])))
        if len(fs) > 1:
            multi += 1
            vistos = set()
            for i, f in enumerate(fs):
                c = curso_de_fila[id(f)]
                if c and c.id not in vistos:
                    vistos.add(c.id)
                    db.add(AlumnoCurso(alumno_id=a.id, curso_id=c.id, anio=ANIO, mes_inicio=c.mes_inicio, principal=i == 0))

    # ------------------------------------------------------------ pagos
    n_pagos, total = 0, Decimal("0")
    for dni, (a, fs) in alumnos.items():
        for f in fs:
            c = curso_de_fila[id(f)]
            for m in range(12):
                imp, fe = f[8 + 2 * m], a_fecha(f[7 + 2 * m])
                if not isinstance(imp, (int, float)) or imp == 0:
                    continue
                fe = fe or date(ANIO, m + 1, 1)
                efectivo = txt(f[54 + m]).lower() == "efectivo"
                h = hashlib.sha1(f"planilla|{dni}|{txt(f[5])}|{m + 1}|{imp}|{fe}".encode()).hexdigest()[:40]
                db.add(PagoCuota(alumno_id=a.id, cliente_siro=None, via="efectivo" if efectivo else "planilla",
                                 importe=Decimal(str(imp)), fecha=fe, fecha_pago=None, canal="Efectivo" if efectivo else None,
                                 mes=fe.month, anio=fe.year, origen="Planilla ALUMNOS 2026", hash=h,
                                 curso_id=c.id if (len(fs) > 1 and c) else None))
                n_pagos += 1
                total += Decimal(str(imp))
    print(f"Alumnos: {len(alumnos)} ({multi} con más de un curso, {sum(1 for a, _ in alumnos.values() if not a.activo)} de baja, "
          f"{sum(1 for a, _ in alumnos.values() if a.legajo.startswith('D'))} sin legajo)")
    print(f"Pagos: {n_pagos} por ${total:,.0f}")
    con_mail = [a for a, _ in alumnos.values() if a.email]
    print(f"Correos para el portal de familias: {len(con_mail)} alumnos con correo válido "
          f"({len({e for a in con_mail for e in a.email.split('; ')})} correos distintos); "
          f"{sum(1 for a, fs in alumnos.values() if txt(fs[0][68]) and not a.email)} con correo ilegible")
    if sin_curso:
        print("SIN CURSO CONFIGURADO:", sin_curso[:10])

    # ------------------------------------------------------------ control horario
    if "--sin-horario" not in sys.argv and "Registro Horario" in wb.sheetnames:
        ws = wb["Registro Horario"]
        filas_h = list(ws.iter_rows(values_only=True))
        titulo = txt(filas_h[0][0]) if filas_h and filas_h[0] else ""
        nombre_emp = "Ana Maria Peini" if "PEINI" in titulo.upper() else (titulo.split("-")[-1].strip().title() or "Empleado")
        emp = db.query(Empleado).filter(Empleado.nombre == nombre_emp).first() or Empleado(
            nombre=nombre_emp, puesto="Administrativa Contable", horas_diarias=Decimal(8), activo=True)
        db.add(emp)
        db.flush()
        regs = 0
        for r in filas_h:
            fe = a_fecha(r[0]) if r else None
            if not fe:
                continue
            tipo = TIPO_DIA.get(txt(r[7]).upper(), "Normal")
            entrada, salida = a_hora(r[4]), a_hora(r[5])
            if tipo == "Normal" and not (entrada and salida):
                continue  # día sin fichada
            db.add(RegistroHorario(empleado_id=emp.id, fecha=fe, tipo=tipo, entrada=entrada, salida=salida))
            regs += 1
        liq = 0
        if "PANEL DE HORAS" in wb.sheetnames:
            meses = {m: i + 1 for i, m in enumerate(["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
                                                       "septiembre", "octubre", "noviembre", "diciembre"])}
            for r in wb["PANEL DE HORAS"].iter_rows(values_only=True):
                if r and txt(r[0]).lower() in meses and isinstance(r[6], (int, float)) and r[6] > 0:
                    db.add(LiquidacionHoras(empleado_id=emp.id, anio=ANIO, mes=meses[txt(r[0]).lower()], horas=Decimal(str(r[6]))))
                    liq += 1
        print(f"Horario: {nombre_emp}, {regs} registros, {liq} liquidaciones de horas")

    db.commit()
    db.close()
    print("Importación terminada.")


if __name__ == "__main__":
    main()
