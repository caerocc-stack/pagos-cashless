"""
Importación masiva de facturas desde una planilla Excel (o CSV).

Flujo: plantilla -> previsualizar (valida, normaliza y detecta duplicados, NO guarda)
-> confirmar (guarda las filas elegidas; las pagadas en efectivo pueden ir al Libro de caja).
Las pagadas por transferencia o débito quedan pendientes: se concilian con el extracto.
"""
import io
import re
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.gasto import Gasto, CATEGORIAS, TIPOS_COMPROBANTE
from app.models.proveedor import Proveedor
from app.models.caja import MovimientoCaja
from app.seguridad import auditar

router = APIRouter(prefix="/api/importar-gastos", tags=["Importar gastos"])

COLUMNAS = [
    ("fecha", "Fecha del comprobante (dd/mm/aaaa)"),
    ("tipo", "Tipo (Factura A/B/C, Ticket, Recibo, Nota de Crédito…)"),
    ("punto_venta", "Punto de venta"),
    ("numero", "Número"),
    ("cuit", "CUIT del proveedor"),
    ("razon_social", "Razón social"),
    ("importe", "Importe total (IVA incluido)"),
    ("concepto", "Concepto (qué se compró)"),
    ("rubro", "Rubro del proveedor"),
    ("categoria", "Categoría"),
    ("destino", "Destino en la escuela"),
    ("medio_pago", "Medio de pago (Efectivo / Transferencia / Débito / Pendiente)"),
    ("fecha_pago", "Fecha de pago (dd/mm/aaaa)"),
    ("notas", "Notas"),
]
MEDIOS = ["Efectivo", "Transferencia", "Débito", "Pendiente"]
ALIAS_CATEGORIA = {"gtos. de oficina": "Oficina", "gastos de oficina": "Oficina", "gtos. de la escuela": "Escuela",
                   "gastos de la escuela": "Escuela", "infraestructura": "Infraestructura", "honorarios": "Honorarios"}


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def _fecha(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _txt(v)
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"fecha inválida «{s}» (usá dd/mm/aaaa)")


def _importe(v) -> Decimal:
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = re.sub(r"[^\d,.\-]", "", _txt(v))
    if not s:
        raise ValueError("falta el importe")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or re.search(r"\.\d{3}$", s):
        s = s.replace(".", "")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise ValueError(f"importe inválido «{_txt(v)}»")


def _cuit(v) -> str | None:
    d = re.sub(r"\D", "", _txt(v))
    if not d:
        return None
    if len(d) != 11:
        raise ValueError(f"CUIT inválido «{_txt(v)}» (deben ser 11 dígitos)")
    return d


def _tipo(v) -> str | None:
    s = _txt(v)
    if not s:
        return None
    u = s.upper().replace(".", "").strip()
    if u in ("A", "B", "C", "FA", "FB", "FC") or re.fullmatch(r"FACTURA ?[ABC]", u) or re.fullmatch(r"FC ?[ABC]", u):
        return f"Factura {u[-1]}"
    if u.startswith(("NC", "NOTA DE CRED", "NOTA CRED")):
        return "Nota de Crédito"
    for t in TIPOS_COMPROBANTE:
        if t.upper() == u:
            return t
    return s[:40]


def _categoria(v) -> str | None:
    s = _txt(v)
    if not s:
        return None
    for c in CATEGORIAS:
        if c.lower() == s.lower():
            return c
    return ALIAS_CATEGORIA.get(s.lower(), s[:40])


def _medio(v) -> str:
    s = _txt(v).lower()
    if not s or s.startswith("pend"):
        return "Pendiente"
    if s.startswith("efec") or s == "ef":
        return "Efectivo"
    if "deb" in s or "déb" in s:
        return "Débito" if "transf" not in s else "Transferencia"
    if "transf" in s or s.startswith("tr"):
        return "Transferencia"
    raise ValueError(f"medio de pago no reconocido «{_txt(v)}»")


def _leer_filas(contenido: bytes, nombre: str) -> list[dict]:
    """Devuelve una lista de dicts {clave: valor} usando los encabezados de la planilla."""
    claves = {k for k, _ in COLUMNAS}

    def mapear_encabezado(h):
        h = _txt(h).lower()
        for k, etiqueta in COLUMNAS:
            if h == k or h == etiqueta.lower() or h.startswith(etiqueta.lower().split(" (")[0]):
                return k
        return None

    if nombre.lower().endswith(".csv"):
        texto = contenido.decode("utf-8-sig", errors="replace")
        sep = ";" if texto.count(";") > texto.count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    else:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
        except Exception:
            raise HTTPException(400, "No se pudo leer el archivo. Subí la planilla .xlsx (o un .csv).")
        ws = wb["Facturas"] if "Facturas" in wb.sheetnames else wb.worksheets[0]
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
    # Buscar la fila de encabezados (la primera que tenga 'fecha' e 'importe')
    for i, fila in enumerate(filas[:10]):
        cab = [mapear_encabezado(h) for h in fila]
        if "fecha" in cab and "importe" in cab:
            salida = []
            for n, r in enumerate(filas[i + 1:], start=i + 2):
                if not any(_txt(x) for x in r):
                    continue
                d = {cab[j]: r[j] for j in range(min(len(cab), len(r))) if cab[j] in claves}
                d["_fila"] = n
                salida.append(d)
            return salida
    raise HTTPException(400, "No encuentro los encabezados. Usá la plantilla (columnas Fecha e Importe como mínimo).")


def _clave_dup(cuit, pv, numero, fecha, importe):
    num = re.sub(r"\D", "", numero or "").lstrip("0")
    if cuit and num:
        return ("num", cuit, re.sub(r"\D", "", pv or "").lstrip("0"), num)
    return ("monto", cuit or "", str(fecha), f"{Decimal(str(importe)):.2f}")


@router.get("/plantilla")
def plantilla():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.append([e for _, e in COLUMNAS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="15233B")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.append(["15/09/2026", "Factura B", "0003", "00012345", "30711050767", "DIAMOND SYSTEM SA", 52360.5,
               "Tóner impresora", "Informática", "Oficina", "Secretaría", "Transferencia", "20/09/2026", ""])
    anchos = [14, 16, 10, 12, 14, 30, 14, 30, 18, 15, 22, 16, 14, 30]
    for i, w in enumerate(anchos):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    dv_cat = DataValidation(type="list", formula1='"' + ",".join(CATEGORIAS) + '"', allow_blank=True)
    dv_med = DataValidation(type="list", formula1='"' + ",".join(MEDIOS) + '"', allow_blank=True)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_med)
    dv_cat.add("J2:J2000")
    dv_med.add("L2:L2000")
    ins = wb.create_sheet("Instrucciones")
    for linea in [
        "Una fila por comprobante. La fila 2 es un ejemplo: borrala o reemplazala.",
        "Obligatorio: Fecha e Importe. Muy recomendado: CUIT, Razón social y Número (evita duplicados y permite conciliar).",
        "Importe: total con IVA. Se aceptan 1234,56 o 1.234,56. Las Notas de Crédito se guardan en negativo.",
        "Categoría: " + ", ".join(CATEGORIAS) + ".",
        "Medio de pago: Efectivo / Transferencia / Débito / Pendiente.",
        "  - Transferencia y Débito quedan pendientes hasta conciliar con el extracto del banco.",
        "  - Efectivo puede registrarse automáticamente como egreso en el Libro de caja.",
        "Si el proveedor no existe (por CUIT), se da de alta solo con la razón social y el rubro.",
    ]:
        ins.append([linea])
    ins.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="plantilla_facturas_APAI.xlsx"'})


@router.post("/previsualizar")
async def previsualizar(archivo: UploadFile = File(...), db: Session = Depends(get_db)):
    """Lee y valida la planilla sin guardar nada: devuelve cada fila normalizada con su estado."""
    filas = _leer_filas(await archivo.read(), archivo.filename or "")
    if not filas:
        raise HTTPException(400, "La planilla no tiene filas con datos.")
    existentes = {_clave_dup(g.cuit and re.sub(r"\D", "", g.cuit), g.punto_venta, g.numero, g.fecha, Decimal(str(g.importe)))
                  for g in db.query(Gasto).all()}
    provs = {p.cuit: p for p in db.query(Proveedor).all()}
    vistas, salida = set(), []
    for f in filas:
        r = {"fila": f["_fila"], "errores": [], "avisos": []}
        try:
            r["fecha"] = _fecha(f.get("fecha"))
            if not r["fecha"]:
                r["errores"].append("falta la fecha")
        except ValueError as e:
            r["errores"].append(str(e))
            r["fecha"] = None
        try:
            r["importe"] = _importe(f.get("importe"))
        except ValueError as e:
            r["errores"].append(str(e))
            r["importe"] = None
        try:
            r["cuit"] = _cuit(f.get("cuit"))
        except ValueError as e:
            r["errores"].append(str(e))
            r["cuit"] = None
        try:
            r["medio_pago"] = _medio(f.get("medio_pago"))
        except ValueError as e:
            r["errores"].append(str(e))
            r["medio_pago"] = "Pendiente"
        try:
            r["fecha_pago"] = _fecha(f.get("fecha_pago"))
        except ValueError as e:
            r["errores"].append("fecha de pago: " + str(e))
            r["fecha_pago"] = None
        r["tipo"] = _tipo(f.get("tipo"))
        r["punto_venta"] = _txt(f.get("punto_venta")).zfill(4) if _txt(f.get("punto_venta")).isdigit() else (_txt(f.get("punto_venta")) or None)
        r["numero"] = _txt(f.get("numero")).zfill(8) if _txt(f.get("numero")).isdigit() else (_txt(f.get("numero")) or None)
        r["razon_social"] = _txt(f.get("razon_social"))[:160] or None
        r["concepto"] = _txt(f.get("concepto"))[:200] or None
        r["rubro"] = _txt(f.get("rubro"))[:80] or None
        r["categoria"] = _categoria(f.get("categoria"))
        r["destino"] = _txt(f.get("destino"))[:160] or None
        r["notas"] = _txt(f.get("notas"))[:400] or None
        if r["tipo"] == "Nota de Crédito" and r["importe"] is not None and r["importe"] > 0:
            r["importe"] = -r["importe"]
        if r["categoria"] and r["categoria"] not in CATEGORIAS:
            r["avisos"].append(f"categoría «{r['categoria']}» no es de la lista: se guarda igual")
        p = provs.get(r["cuit"]) if r["cuit"] else None
        if p:
            r["proveedor"] = p.razon_social
            r["razon_social"] = r["razon_social"] or p.razon_social
            r["rubro"] = r["rubro"] or p.rubro
        elif r["cuit"] and r["razon_social"]:
            r["avisos"].append("proveedor nuevo: se da de alta")
        elif not r["cuit"]:
            r["avisos"].append("sin CUIT: no se podrá conciliar automáticamente")
        if r["medio_pago"] == "Efectivo" and not r["fecha_pago"]:
            r["fecha_pago"] = r["fecha"]
        if not r["errores"]:
            k = _clave_dup(r["cuit"], r["punto_venta"], r["numero"], r["fecha"], r["importe"])
            if k in existentes:
                r["estado"] = "duplicado"
            elif k in vistas:
                r["estado"] = "repetido"
            else:
                r["estado"] = "nuevo"
                vistas.add(k)
        else:
            r["estado"] = "error"
        for c in ("fecha", "fecha_pago"):
            r[c] = str(r[c]) if r[c] else None
        r["importe"] = float(r["importe"]) if r["importe"] is not None else None
        salida.append(r)
    resumen = {e: sum(1 for x in salida if x["estado"] == e) for e in ("nuevo", "duplicado", "repetido", "error")}
    resumen["total_nuevos"] = round(sum(x["importe"] for x in salida if x["estado"] == "nuevo"), 2)
    return {"filas": salida, "resumen": resumen}


class FilaImport(BaseModel):
    fecha: date
    tipo: str | None = None
    punto_venta: str | None = None
    numero: str | None = None
    cuit: str | None = None
    razon_social: str | None = None
    importe: Decimal
    concepto: str | None = None
    rubro: str | None = None
    categoria: str | None = None
    destino: str | None = None
    medio_pago: str = "Pendiente"
    fecha_pago: date | None = None
    notas: str | None = None


class Confirmacion(BaseModel):
    filas: list[FilaImport]
    efectivo_a_caja: bool = True


@router.post("/confirmar")
def confirmar(data: Confirmacion, request: Request, db: Session = Depends(get_db)):
    """Guarda las filas elegidas (vuelve a controlar duplicados por si cambió algo)."""
    existentes = {_clave_dup(g.cuit and re.sub(r"\D", "", g.cuit), g.punto_venta, g.numero, g.fecha, Decimal(str(g.importe)))
                  for g in db.query(Gasto).all()}
    provs = {p.cuit: p for p in db.query(Proveedor).all()}
    creados, omitidos, a_caja, prov_nuevos, total = 0, 0, 0, 0, Decimal("0")
    for f in data.filas:
        cuit = re.sub(r"\D", "", f.cuit or "") or None
        k = _clave_dup(cuit, f.punto_venta, f.numero, f.fecha, f.importe)
        if k in existentes:
            omitidos += 1
            continue
        existentes.add(k)
        p = provs.get(cuit) if cuit else None
        if not p and cuit and f.razon_social:
            p = Proveedor(cuit=cuit, razon_social=f.razon_social, rubro=f.rubro)
            db.add(p)
            db.flush()
            provs[cuit] = p
            prov_nuevos += 1
        notas = f.notas
        if f.medio_pago in ("Transferencia", "Débito"):
            # Queda pendiente: el pago se confirma al conciliar con el extracto
            notas = "; ".join(x for x in [f"Pago anotado: {f.medio_pago}" + (f" {f.fecha_pago:%d/%m/%Y}" if f.fecha_pago else ""), notas] if x)
        g = Gasto(fecha=f.fecha, tipo_comprobante=f.tipo, punto_venta=f.punto_venta, numero=f.numero,
                  proveedor_id=p.id if p else None, razon_social=f.razon_social or (p.razon_social if p else None),
                  cuit=cuit, importe=f.importe, concepto=f.concepto, rubro=f.rubro or (p.rubro if p else None),
                  categoria=f.categoria, destino=f.destino, notas=(notas or None) and notas[:400])
        db.add(g)
        db.flush()
        if f.medio_pago == "Efectivo":
            g.forma_pago = "Efectivo"
            g.fecha_pago = f.fecha_pago or f.fecha
            if data.efectivo_a_caja and f.importe > 0:
                c = MovimientoCaja(fecha=g.fecha_pago, concepto=f"Pago {g.razon_social or 'proveedor'}"[:200],
                                   tipo_movimiento="Pago proveedor", tipo="egreso", monto=abs(f.importe))
                db.add(c)
                db.flush()
                g.caja_mov_id = c.id
                a_caja += 1
        creados += 1
        total += f.importe
    auditar(db, "importacion_facturas", None, f"{creados} facturas importadas, {omitidos} duplicadas omitidas", request)
    db.commit()
    return {"ok": True, "creados": creados, "omitidos": omitidos, "en_caja": a_caja, "proveedores_nuevos": prov_nuevos,
            "total": float(total),
            "mensaje": f"Se importaron {creados} comprobantes" + (f" ({omitidos} ya existían)" if omitidos else "") +
                       (f", {a_caja} pagos en efectivo registrados en Caja" if a_caja else "") +
                       (f" y {prov_nuevos} proveedores nuevos" if prov_nuevos else "") + "."}
