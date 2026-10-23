"""
Planilla Excel para cargar los cupones de la cuota en SIRO (modelo "carga de cupones SIRO").

Columnas: Nro Cliente | 1º Vto | 1º Monto | 2º Vto | 2º Monto | 3º Vto | 3º Monto.
Solo se completan el número de cliente, el primer vencimiento y el primer monto.

Número de cliente, tal cual, sin agregarle dígitos:
- Aeronavegantes y exámenes: el DNI.
- Ciclo Básico, Aviónica y Mecánica: el legajo.
"""
import io
import re
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.auth import requiere_admin
from app.cuenta_alumno import estado_cuenta
from app.database import get_db
from app.models.alumno import Alumno
from app.routers.cuotas import cursos_vigentes, extras_por_alumno, pagos_del_anio, planes_de
from app.seguridad import auditar
from app.tz import ahora_ar

router = APIRouter(prefix="/api/cupones-siro/planilla", tags=["Cupones SIRO"])

AREAS_CON_DNI = {"aeronavegantes", "examenes", "exámenes"}
ORDEN_AREAS = ["Ciclo Basico", "Avionica", "Mecanica", "Aeronavegantes", "Examenes"]
ENCABEZADOS = ["Nro Cliente", "1º Vto", "1º Monto", "2º Vto", "2º Monto", "3º Vto", "3º Monto"]


class CursoPlanilla(BaseModel):
    curso_id: int
    importe: Decimal | None = Field(default=None, gt=0, le=10_000_000)


class Pedido(BaseModel):
    vencimiento: date
    cursos: list[CursoPlanilla] = Field(min_length=1, max_length=50)
    anio: int | None = None
    excluir_pagados: bool = True
    # alumnos que no reciben cupón ("curso_id:alumno_id"), p. ej. en un examen solo se manda a los que rinden
    excluidos: list[str] = Field(default=[], max_length=5000)


def usa_dni(area: str | None) -> bool:
    return (area or "").strip().lower() in AREAS_CON_DNI


def _exento(a: Alumno) -> bool:
    return (a.condicion or "").strip().lower() in ("becado", "exento") or bool(a.cuota_excluir)


def _anio(p: int | None) -> int:
    return p or ahora_ar().year


def _alumnos_con_planes(db: Session, anio: int):
    cursos = cursos_vigentes(db, anio)
    extras = extras_por_alumno(db, anio)
    for a in db.query(Alumno).filter(Alumno.activo == True).order_by(Alumno.apellido, Alumno.nombre).all():  # noqa: E712
        yield a, planes_de(a, cursos, extras.get(a.id, []))


def _ordenar_cursos(cursos):
    orden = {x.lower(): i for i, x in enumerate(ORDEN_AREAS)}
    return sorted(cursos, key=lambda c: (orden.get((c.area or "").lower(), 99), c.nombre))


def entradas(db: Session, anio: int, curso_ids: set[int] | None = None) -> list[dict]:
    """Un renglón por alumno y curso, con su número de cliente y si recibe cupón o por qué no
    (becado, año pagado, sin número)."""
    cursos = {c.id: c for c in cursos_vigentes(db, anio)}
    pagos = pagos_del_anio(db, anio)
    salida = []
    for a, planes in _alumnos_con_planes(db, anio):
        elegidos = [pl for pl in planes if pl.get("curso_id") in cursos and (curso_ids is None or pl["curso_id"] in curso_ids)]
        if not elegidos:
            continue
        exento = _exento(a)
        completos = set()
        if not exento:
            info = estado_cuenta(a, planes, pagos.get(a.id, []), anio, hoy=ahora_ar().date())
            completos = {c["clave"] for c in info.get("cursos", []) if c.get("completo")}
        for pl in elegidos:
            curso = cursos[pl["curso_id"]]
            dni = usa_dni(curso.area)
            numero = re.sub(r"\D", "", a.dni or "") if dni else (a.legajo or "").strip()
            if exento:
                estado, motivo = "becado", (a.condicion or "Exento").capitalize()
            elif not numero.isdigit():
                estado = "sin_numero"
                motivo = "No tiene DNI cargado" if dni else f"No tiene legajo numérico (figura «{a.legajo}»)"
            elif pl["clave"] in completos:
                estado, motivo = "pagado", "Ya pagó el año completo"
            else:
                estado, motivo = "ok", None
            salida.append({
                "curso_id": curso.id, "alumno_id": a.id, "alumno": f"{a.apellido}, {a.nombre}",
                "etiqueta": pl.get("etiqueta") or (a.curso if pl.get("principal") else curso.nombre),
                "cliente": numero if numero.isdigit() else None, "tipo": "DNI" if dni else "Legajo",
                "personal": float(a.cuota_personalizada) if pl.get("principal") and a.cuota_personalizada else None,
                "estado": estado, "motivo": motivo})
    return salida


@router.get("/cursos")
def cursos_para_planilla(anio: int | None = None, db: Session = Depends(get_db), _=Depends(requiere_admin)):
    """Cursos del año (importe, qué número de cliente usan) con todos sus alumnos y si reciben cupón.
    La pantalla arma la vista previa con esto, sin volver a consultar."""
    anio = _anio(anio)
    por_curso: dict[int, list] = {}
    for e in entradas(db, anio):
        por_curso.setdefault(e.pop("curso_id"), []).append(e)
    return [{"id": c.id, "nombre": c.nombre, "area": c.area, "importe": float(c.cuota or 0), "por_examen": bool(c.por_examen),
             "cliente": "DNI" if usa_dni(c.area) else "Legajo", "alumnos": por_curso.get(c.id, [])}
            for c in _ordenar_cursos(cursos_vigentes(db, anio))]


def armar_filas(db: Session, data: Pedido) -> dict:
    anio = _anio(data.anio)
    cursos = {c.id: c for c in cursos_vigentes(db, anio)}
    pedidos = {}
    for x in data.cursos:
        if x.curso_id not in cursos:
            raise HTTPException(400, f"El curso {x.curso_id} no existe o no está activo en {anio}")
        pedidos[x.curso_id] = Decimal(str(x.importe)) if x.importe else Decimal(cursos[x.curso_id].cuota or 0)

    filas: dict[str, dict] = {}
    omitidos = {"becados": 0, "pagados": 0, "elegidos": 0}
    sin_numero = []
    excluidos = set(data.excluidos)
    for e in entradas(db, anio, set(pedidos)):
        if f"{e['curso_id']}:{e['alumno_id']}" in excluidos:
            omitidos["elegidos"] += 1
            continue
        if e["estado"] == "becado":
            omitidos["becados"] += 1
            continue
        if e["estado"] == "sin_numero":
            sin_numero.append({"alumno": e["alumno"], "curso": e["etiqueta"], "motivo": e["motivo"]})
            continue
        if e["estado"] == "pagado" and data.excluir_pagados:
            omitidos["pagados"] += 1
            continue
        importe = Decimal(str(e["personal"])) if e["personal"] else pedidos[e["curso_id"]]
        f = filas.get(e["cliente"])
        if f:  # mismo número de cliente en dos cursos elegidos: un solo cupón con la suma
            f["importe"] += importe
            f["cursos"].append(e["etiqueta"])
            f["por_curso"].append((e["curso_id"], importe))
        else:
            filas[e["cliente"]] = {"cliente": e["cliente"], "tipo": e["tipo"], "alumno_id": e["alumno_id"],
                                   "alumno": e["alumno"], "cursos": [e["etiqueta"]], "importe": importe,
                                   "personalizado": bool(e["personal"]), "por_curso": [(e["curso_id"], importe)]}
    lista = sorted(filas.values(), key=lambda f: (f["cursos"][0], f["alumno"]))
    return {"filas": lista, "omitidos": omitidos, "sin_numero": sin_numero, "anio": anio,
            "total": sum((f["importe"] for f in lista), Decimal(0))}


@router.post("/vista")
def vista(data: Pedido, db: Session = Depends(get_db), _=Depends(requiere_admin)):
    r = armar_filas(db, data)
    return {**r, "total": float(r["total"]),
            "filas": [{**{k: v for k, v in f.items() if k != "por_curso"}, "importe": float(f["importe"]),
                       "cursos": " + ".join(f["cursos"])} for f in r["filas"]]}


def registrar_cupones(db: Session, cupones: list[tuple], vencimiento: date, origen: str, lote: str) -> int:
    """Guarda (alumno, curso, cliente, monto) como cupón generado con ese vencimiento; si ya estaba, lo actualiza."""
    from app.models.curso import CuponGenerado
    ya = {(c.alumno_id, c.curso_id): c for c in db.query(CuponGenerado).filter(CuponGenerado.vencimiento == vencimiento)}
    nuevos = 0
    for alumno_id, curso_id, cliente, monto in cupones:
        c = ya.get((alumno_id, curso_id))
        if c:
            c.monto, c.cliente, c.lote, c.origen = Decimal(str(monto)), cliente, lote[:120], origen
        else:
            c = CuponGenerado(alumno_id=alumno_id, curso_id=curso_id, cliente=cliente, vencimiento=vencimiento,
                              monto=Decimal(str(monto)), origen=origen, lote=lote[:120])
            db.add(c)
            ya[(alumno_id, curso_id)] = c
            nuevos += 1
    return nuevos


def generar_excel(filas: list[dict], vencimiento: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    borde = Border(*(Side(style="thin", color="000000"),) * 4)
    gris = PatternFill("solid", fgColor="A6A6A6")
    for i, t in enumerate(ENCABEZADOS, 1):
        c = ws.cell(row=1, column=i, value=t)
        c.font, c.fill, c.border, c.alignment = Font(bold=True), gris, borde, Alignment(horizontal="center")
        c.number_format = "mm-dd-yy" if i == 2 else ("0.00" if i in (3, 5, 7) else "General")
    for r, f in enumerate(filas, 2):
        valores = [int(f["cliente"]), vencimiento, float(f["importe"]), None, None, None, None]
        for i, v in enumerate(valores, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = borde
        ws.cell(row=r, column=2).number_format = "mm-dd-yy"
        ws.cell(row=r, column=3).number_format = "0.00"
    for letra, ancho in zip("ABCDEFG", (14, 12, 12, 10, 10, 10, 10)):
        ws.column_dimensions[letra].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/excel")
def excel(data: Pedido, request: Request, db: Session = Depends(get_db), admin=Depends(requiere_admin)):
    r = armar_filas(db, data)
    if not r["filas"]:
        raise HTTPException(400, "No hay alumnos para generar cupones con esos cursos")
    nombres = sorted({c for f in r["filas"] for c in f["cursos"]})
    auditar(db, "planilla_siro", admin["sub"],
            f"{len(r['filas'])} cupones, vence {data.vencimiento:%d/%m/%Y}, total ${r['total']:,.2f}: {', '.join(nombres)[:300]}", request)
    db.commit()
    nombre = f"cupones_siro_{data.vencimiento:%Y-%m-%d}.xlsx"
    return StreamingResponse(io.BytesIO(generar_excel(r["filas"], data.vencimiento)),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})
