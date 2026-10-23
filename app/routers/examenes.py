"""
Cursos que se rinden por examen (RAAC 147, Pericia): instancias del año y alumnos que pagaron cada una.

Cada pago de estos cursos es la inscripción a un examen. Va a la primera instancia desde el mes en que
se pagó (fecha de pago), salvo que se la haya movido a mano a otra (pago.instancia_id).

Los inscriptos de una instancia son los alumnos a los que se les mandó cupón para ese examen: los que se
eligieron al generar la planilla de SIRO (o los de una planilla ya subida a SIRO que se registró acá).
El cupón va a la primera instancia desde el mes de su 1º vencimiento.
"""
import io
import re
from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.alumno import Alumno
from app.models.curso import Curso, InstanciaExamen, CuponGenerado
from app.models.pago_cuota import PagoCuota
from app.cuenta_alumno import MESES, instancia_de, separar_pagos_examen, fecha_imputacion
from app.routers.cuotas import cursos_vigentes, extras_por_alumno, pagos_del_anio, planes_de
from app.tz import ahora_ar

router = APIRouter(prefix="/api/examenes", tags=["Exámenes"])


def _curso_examen(db: Session, curso_id: int) -> Curso:
    c = db.get(Curso, curso_id)
    if not c or not c.por_examen:
        raise HTTPException(404, "Curso por examen no encontrado")
    return c


def _inst_dto(i: InstanciaExamen) -> dict:
    return {"id": i.id, "anio": i.anio, "mes": i.mes, "fecha": str(i.fecha) if i.fecha else None,
            "descripcion": i.descripcion, "nombre": f"{MESES[i.mes - 1].capitalize()} {i.anio}"}


@router.get("/cursos")
def cursos(db: Session = Depends(get_db)):
    return [{"id": c.id, "nombre": c.nombre, "cuota": float(c.cuota or 0), "anio": c.anio, "activo": c.activo}
            for c in db.query(Curso).filter(Curso.por_examen.is_(True)).order_by(Curso.nombre).all()]


@router.get("/{curso_id}")
def detalle(curso_id: int, anio: int | None = None, db: Session = Depends(get_db)):
    """Instancias del año y, en cada una, los alumnos que pagaron (con la fecha de proceso del pago)."""
    anio = anio or ahora_ar().year
    curso = _curso_examen(db, curso_id)
    instancias = (db.query(InstanciaExamen).filter(InstanciaExamen.curso_id == curso_id, InstanciaExamen.anio == anio)
                  .order_by(InstanciaExamen.mes).all())
    # el mismo cálculo que la cuenta del alumno: qué pagos son de este curso por examen
    cursos_cfg = cursos_vigentes(db, anio)
    if not any(c.id == curso_id for c in cursos_cfg):
        cursos_cfg = cursos_cfg + [curso]
    extras = extras_por_alumno(db, anio)
    pagos = pagos_del_anio(db, anio)
    filas = []
    for a in db.query(Alumno).filter(Alumno.id.in_(list(pagos.keys()) or [0])).all():
        planes = planes_de(a, cursos_cfg, extras.get(a.id, []))
        claves = {pl["clave"] for pl in planes if pl.get("curso_id") == curso_id}
        if not claves:
            continue
        _, por_plan = separar_pagos_examen(planes, pagos.get(a.id, []))
        for clave in claves:
            for p in por_plan.get(clave, []):
                inst = instancia_de(p, instancias)
                filas.append({"pago_id": p.id, "alumno_id": a.id, "legajo": a.legajo, "dni": a.dni, "apellido": a.apellido,
                              "nombre": a.nombre, "curso": a.curso, "importe": float(p.importe), "fecha": str(p.fecha),
                              "fecha_pago": str(p.fecha_pago) if p.fecha_pago else None, "canal": p.canal, "via": p.via,
                              "instancia_id": inst.id if inst else None, "manual": bool(p.instancia_id),
                              "mes_pago": fecha_imputacion(p).month})
    filas.sort(key=lambda f: (f["apellido"], f["nombre"], f["fecha"]))
    # inscriptos de cada instancia: a quiénes se les mandó cupón para ese examen
    insc_por_inst: dict[int, dict] = {}
    cupones_sin_instancia = 0
    for c, a in (db.query(CuponGenerado, Alumno).join(Alumno, Alumno.id == CuponGenerado.alumno_id)
                 .filter(CuponGenerado.curso_id == curso_id, CuponGenerado.vencimiento >= date(anio - 1, 12, 1),
                         CuponGenerado.vencimiento <= date(anio, 12, 31))
                 .order_by(CuponGenerado.vencimiento).all()):
        inst = instancia_de_cupon(c.vencimiento, instancias)
        if not inst:
            cupones_sin_instancia += 1
            continue
        insc_por_inst.setdefault(inst.id, {})[a.id] = (a, c)  # si tiene dos cupones para la misma, vale el último
    otros: dict[int, list] = {}  # pagos del alumno en el año (para ver si pagó en otra instancia)
    for f in filas:
        otros.setdefault(f["alumno_id"], []).append(f)
    salida = []
    for i in instancias:
        de_esta = [f for f in filas if f["instancia_id"] == i.id]
        pagaron = {f["alumno_id"] for f in de_esta}
        inscriptos = sorted(insc_por_inst.get(i.id, {}).values(), key=lambda x: (x[0].apellido, x[0].nombre))
        ids_insc = {a.id for a, _ in inscriptos}
        for f in de_esta:
            f["inscripto"] = f["alumno_id"] in ids_insc if inscriptos else None
        salida.append({**_inst_dto(i), "cantidad": len(de_esta), "alumnos": len(pagaron),
                       "total": sum(f["importe"] for f in de_esta), "pagos": de_esta,
                       "inscriptos": [{"alumno_id": a.id, "apellido": a.apellido, "nombre": a.nombre, "dni": a.dni,
                                       "legajo": a.legajo, "curso": a.curso, "pago": a.id in pagaron,
                                       "cupon_id": c.id, "vencimiento": str(c.vencimiento), "monto": float(c.monto or 0),
                                       "origen": c.origen,
                                       "pago_fecha": next((f["fecha"] for f in de_esta if f["alumno_id"] == a.id), None),
                                       "otras_instancias": sorted({f["instancia_id"] for f in otros.get(a.id, []) if f["instancia_id"] != i.id and f["instancia_id"]})}
                                      for a, c in inscriptos],
                       "no_pagaron": sum(1 for a, _ in inscriptos if a.id not in pagaron)})
    sin = [f for f in filas if f["instancia_id"] is None]
    return {"curso": {"id": curso.id, "nombre": curso.nombre, "cuota": float(curso.cuota or 0)}, "anio": anio,
            "instancias": salida, "sin_instancia": sin,
            "meses_sin_instancia": sorted({f["mes_pago"] for f in sin}), "cupones_sin_instancia": cupones_sin_instancia,
            "total_pagos": len(filas), "total": sum(f["importe"] for f in filas)}


def instancia_de_cupon(vencimiento: date, instancias: list):
    """Primera instancia desde el mes del 1º vencimiento del cupón (el mismo criterio que los pagos)."""
    return next((i for i in sorted(instancias, key=lambda i: (i.anio, i.mes))
                 if (i.anio, i.mes) >= (vencimiento.year, vencimiento.month)), None)


class InstanciaIn(BaseModel):
    anio: int
    mes: int
    fecha: date | None = None
    descripcion: str | None = None


def _validar(db: Session, curso_id: int, data: InstanciaIn, excluir: int | None = None):
    if not 1 <= data.mes <= 12:
        raise HTTPException(400, "Mes inválido")
    if data.fecha and (data.fecha.year, data.fecha.month) != (data.anio, data.mes):
        raise HTTPException(400, "El día del examen tiene que ser del mes elegido.")
    otra = (db.query(InstanciaExamen).filter(InstanciaExamen.curso_id == curso_id, InstanciaExamen.anio == data.anio,
                                             InstanciaExamen.mes == data.mes, InstanciaExamen.id != (excluir or 0)).first())
    if otra:
        raise HTTPException(409, f"Ya hay una instancia en {MESES[data.mes - 1]} {data.anio}.")


@router.post("/{curso_id}/instancias")
def crear_instancia(curso_id: int, data: InstanciaIn, db: Session = Depends(get_db)):
    _curso_examen(db, curso_id)
    _validar(db, curso_id, data)
    i = InstanciaExamen(curso_id=curso_id, anio=data.anio, mes=data.mes, fecha=data.fecha,
                        descripcion=(data.descripcion or "").strip()[:120] or None)
    db.add(i)
    db.commit()
    return _inst_dto(i)


@router.put("/instancias/{inst_id}")
def editar_instancia(inst_id: int, data: InstanciaIn, db: Session = Depends(get_db)):
    i = db.get(InstanciaExamen, inst_id)
    if not i:
        raise HTTPException(404, "Instancia no encontrada")
    _validar(db, i.curso_id, data, inst_id)
    i.anio, i.mes, i.fecha = data.anio, data.mes, data.fecha
    i.descripcion = (data.descripcion or "").strip()[:120] or None
    db.commit()
    return _inst_dto(i)


@router.delete("/instancias/{inst_id}")
def borrar_instancia(inst_id: int, db: Session = Depends(get_db)):
    i = db.get(InstanciaExamen, inst_id)
    if not i:
        raise HTTPException(404, "Instancia no encontrada")
    db.query(PagoCuota).filter(PagoCuota.instancia_id == inst_id).update({PagoCuota.instancia_id: None})
    db.delete(i)
    db.commit()
    return {"ok": True, "mensaje": "Instancia eliminada: sus pagos vuelven a asignarse por la fecha de pago"}


class MoverPago(BaseModel):
    instancia_id: int | None = None  # None: vuelve a la asignación automática por fecha


@router.put("/pagos/{pago_id}/instancia")
def mover_pago(pago_id: int, data: MoverPago, db: Session = Depends(get_db)):
    p = db.get(PagoCuota, pago_id)
    if not p:
        raise HTTPException(404, "Pago no encontrado")
    if data.instancia_id and not db.get(InstanciaExamen, data.instancia_id):
        raise HTTPException(404, "Instancia no encontrada")
    p.instancia_id = data.instancia_id
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- cupones ya enviados (inscriptos)

def _alumnos_del_curso(db: Session, curso: Curso) -> dict[int, Alumno]:
    anio = curso.anio or ahora_ar().year
    cfg = cursos_vigentes(db, anio)
    if not any(c.id == curso.id for c in cfg):
        cfg = cfg + [curso]
    extras = extras_por_alumno(db, anio)
    return {a.id: a for a in db.query(Alumno).all()
            if any(pl.get("curso_id") == curso.id for pl in planes_de(a, cfg, extras.get(a.id, [])))}


def _buscar(alumnos: dict[int, Alumno], cliente: str) -> Alumno | None:
    c = re.sub(r"\D", "", cliente or "").lstrip("0")
    if not c:
        return None
    return (next((a for a in alumnos.values() if re.sub(r"\D", "", a.dni or "").lstrip("0") == c), None)
            or next((a for a in alumnos.values() if re.sub(r"\D", "", a.legajo or "").lstrip("0") == c), None))


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\s*$", str(v or ""))
    if not m:
        return None
    d, mes, a = map(int, m.groups())
    try:
        return date(a + 2000 if a < 100 else a, mes, d)
    except ValueError:
        return None


def leer_planilla_siro(contenido: bytes) -> list[tuple[str, date | None, Decimal | None]]:
    """(Nro Cliente, 1º Vto, 1º Monto) de cada renglón de una planilla de carga de cupones de SIRO."""
    from openpyxl import load_workbook
    try:
        ws = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo: tiene que ser la planilla Excel (.xlsx) que se subió a SIRO")
    filas, cols = [], None
    for r in ws.iter_rows(values_only=True):
        textos = [str(x or "").strip().lower() for x in r]
        if cols is None:
            if any("cliente" in t for t in textos):
                cols = {"cliente": next(k for k, t in enumerate(textos) if "cliente" in t),
                        "vto": next((k for k, t in enumerate(textos) if "vto" in t or "venc" in t), None),
                        "monto": next((k for k, t in enumerate(textos) if "monto" in t or "importe" in t), None)}
            continue
        cli = r[cols["cliente"]] if cols["cliente"] < len(r) else None
        if cli is None or str(cli).strip() == "":
            continue
        cli = str(int(cli)) if isinstance(cli, float) else str(cli).strip()
        vto = _fecha(r[cols["vto"]]) if cols["vto"] is not None and cols["vto"] < len(r) else None
        monto = None
        if cols["monto"] is not None and cols["monto"] < len(r) and r[cols["monto"]] not in (None, ""):
            try:
                monto = Decimal(str(r[cols["monto"]]).replace("$", "").strip())
            except Exception:
                monto = None
        filas.append((cli, vto, monto))
    if cols is None:
        raise HTTPException(400, "No encontré la columna «Nro Cliente» en la planilla")
    return filas


def _registrar(db: Session, curso: Curso, filas: list[tuple], vencimiento: date | None, lote: str) -> dict:
    from app.routers.planilla_siro import registrar_cupones
    alumnos = _alumnos_del_curso(db, curso)
    por_vto: dict[date, list] = {}
    no_encontrados, sin_fecha = [], 0
    for cli, vto, monto in filas:
        vto = vto or vencimiento
        if not vto:
            sin_fecha += 1
            continue
        a = _buscar(alumnos, cli)
        if not a:
            no_encontrados.append(cli)
            continue
        por_vto.setdefault(vto, []).append((a.id, curso.id, re.sub(r"\D", "", cli), monto or curso.cuota or 0))
    if sin_fecha and not por_vto:
        raise HTTPException(400, "La planilla no trae el 1º vencimiento: poné la fecha de vencimiento de esos cupones")
    nuevos = sum(registrar_cupones(db, cupones, vto, "importado", lote) for vto, cupones in por_vto.items())
    db.commit()
    return {"ok": True, "registrados": sum(len(x) for x in por_vto.values()), "nuevos": nuevos,
            "vencimientos": sorted(str(v) for v in por_vto), "no_encontrados": no_encontrados}


@router.post("/{curso_id}/cupones/planilla")
async def registrar_planilla(curso_id: int, archivo: UploadFile = File(...), vencimiento: date | None = Form(None),
                             db: Session = Depends(get_db)):
    """Registra los cupones de una planilla que ya se subió a SIRO: esos alumnos quedan inscriptos al examen."""
    curso = _curso_examen(db, curso_id)
    contenido = await archivo.read()
    if len(contenido) > 5_000_000:
        raise HTTPException(413, "El archivo es demasiado grande")
    return _registrar(db, curso, leer_planilla_siro(contenido), vencimiento,
                      f"Planilla subida: {(archivo.filename or 'cupones.xlsx')[:90]}")


class CuponesPegados(BaseModel):
    vencimiento: date
    clientes: list[str] = Field(min_length=1, max_length=2000)  # DNI (o legajo) de cada alumno que recibió cupón
    monto: Decimal | None = Field(default=None, gt=0, le=10_000_000)


@router.post("/{curso_id}/cupones")
def registrar_lista(curso_id: int, data: CuponesPegados, db: Session = Depends(get_db)):
    """Registra cupones ya enviados a partir de una lista de DNI (cuando no se tiene la planilla de SIRO)."""
    curso = _curso_examen(db, curso_id)
    return _registrar(db, curso, [(c, None, data.monto) for c in data.clientes], data.vencimiento,
                      f"Lista cargada el {ahora_ar():%d/%m/%Y}")


@router.delete("/cupones/{cupon_id}")
def quitar_cupon(cupon_id: int, db: Session = Depends(get_db)):
    c = db.get(CuponGenerado, cupon_id)
    if not c:
        raise HTTPException(404, "Cupón no encontrado")
    db.delete(c)
    db.commit()
    return {"ok": True}
