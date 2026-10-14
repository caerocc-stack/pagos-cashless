"""
Excel mensual de gastos para la contadora, con el formato de la planilla modelo de APAI
("Gastos MARZO-2026.xlsx", el mismo que armaba scripts/build_gastos_excel.py):

- TOTALES: "GASTOS", el mes, y CONCEPTO / TOTAL por destino con el TOTAL general (verde).
- Una hoja por destino: Gtos. de Oficina, Infraestructura, Gtos. de la Escuela, Honorarios.
  Columnas COMPROBANTES (fecha, letra, Nº, razón social, C.U.I.T., importe, concepto, rubro) y
  PAGOS (fecha, medio), con el total de la hoja en amarillo. Los totales son fórmulas.
- A REVISAR (solo si hace falta): facturas sin destino (no suman en ninguna hoja) y facturas
  todavía sin revisar (sí suman: se listan para confirmarlas).

Formatos de número portables para que se vea igual en Excel y en Google Sheets.
"""
import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# destino en el sistema -> hoja del Excel
HOJAS = {"Oficina": "Gtos. de Oficina", "Infraestructura": "Infraestructura",
         "Escuela": "Gtos. de la Escuela", "Honorarios": "Honorarios"}
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre",
         "Octubre", "Noviembre", "Diciembre"]

FONT = "Rubik"
HEADER_FONT = Font(name=FONT, size=11, bold=True)
BODY_FONT = Font(name=FONT, size=11)
TITLE_FONT = Font(name="Courier New", size=16, bold=True)
SUBTITLE_FONT = Font(name="Courier New", size=15, bold=True)
TOT_HEADER_FONT = Font(name=FONT, size=12, bold=True)
TOT_ROW_FONT = Font(name=FONT, size=11)
MONEY = r'"$"\ #,##0.00'  # el mismo formato que guarda Excel en la planilla modelo
AMARILLO = PatternFill(fill_type="solid", fgColor="FFFFFF00")
VERDE = PatternFill(fill_type="solid", fgColor="FF92D050")
NARANJA = PatternFill(fill_type="solid", fgColor="FFFCE4D6")
MEDIUM = Side(style="medium", color="FF000000")
THIN = Side(style="thin", color="FF000000")
CAJA = Border(top=MEDIUM, bottom=MEDIUM, left=MEDIUM, right=MEDIUM)
LADOS = Border(left=MEDIUM, right=MEDIUM)
ANCHOS = {"A": 1.71, "B": 12.43, "C": 21.43, "D": 20.0, "E": 42.29, "F": 18.29, "G": 21.14, "H": 56.29,
          "I": 21.14, "J": 9.86, "K": 16.14, "L": 10.71}
ALINEACION = {2: "center", 3: "center", 4: "center", 5: None, 6: "center", 7: "right", 8: "left", 9: "left",
              10: "center", 11: "center"}
ALINEACION_CAB = {4: "center", 5: "center", 10: "center", 11: "center"}


def fecha_corta(d: date | None) -> str | None:
    return d.strftime("%d/%m/%y") if d else None


def numero_comprobante(pv, numero) -> str | None:
    pv, numero = (str(pv or "").strip(), str(numero or "").strip())
    if not numero:
        return pv or None
    if not pv:
        return numero
    pv = pv.zfill(4) if pv.isdigit() else pv  # como figura en la factura: 0012 o 00012
    numero = numero.zfill(8) if numero.isdigit() else numero
    return f"{pv}-{numero}"


def cuit_con_guiones(c) -> str | None:
    d = re.sub(r"\D", "", str(c or ""))
    return f"{d[:2]}-{d[2:10]}-{d[10:]}" if len(d) == 11 else (c or None)


def medio(forma_pago) -> str | None:
    if not forma_pago:
        return None
    return "Transf/Debito" if forma_pago in ("Transferencia", "Débito") else forma_pago


def fila(g) -> dict:
    return {2: fecha_corta(g.fecha), 3: g.tipo_comprobante, 4: numero_comprobante(g.punto_venta, g.numero),
            5: g.razon_social, 6: cuit_con_guiones(g.cuit), 7: float(g.importe),
            8: ((g.concepto or g.destino or "") + (f" (pagado vía {g.pagado_por})" if getattr(g, "pagado_por", None) else "")) or None, 9: g.rubro, 10: fecha_corta(g.fecha_pago), 11: medio(g.forma_pago)}


def _hoja_destino(ws, gastos) -> int:
    ws.cell(row=2, column=2, value="COMPROBANTES")
    ws.cell(row=2, column=10, value="PAGOS")
    for col in (2, 10):
        ws.cell(row=2, column=col).font = HEADER_FONT
        ws.cell(row=2, column=col).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)
    ws.merge_cells(start_row=2, start_column=10, end_row=2, end_column=11)
    ws.row_dimensions[2].height = ws.row_dimensions[3].height = 16.5
    for i, h in enumerate(["FECHA", "LETRA", "Nº ", "RAZÓN SOCIAL", "C.U.I.T.", "IMPORTE", "CONCEPTO", "RUBRO", "FECHA", "MEDIO"]):
        c = ws.cell(row=3, column=2 + i, value=h)
        c.font, c.border = HEADER_FONT, CAJA
        c.alignment = Alignment(horizontal=ALINEACION_CAB.get(2 + i), vertical="center")
    r = 4
    for g in gastos:
        for col, val in fila(g).items():
            c = ws.cell(row=r, column=col, value=val)
            c.font, c.border = BODY_FONT, LADOS
            c.alignment = Alignment(horizontal=ALINEACION.get(col), vertical="center", wrap_text=True)
            if col == 7:
                c.number_format = MONEY
        r += 1
    total = r
    for col in range(2, 12):
        ws.cell(row=total, column=col).border = Border(top=MEDIUM)
    et = ws.cell(row=total, column=6, value="Total")
    suma = ws.cell(row=total, column=7, value=f"=SUM(G4:G{r - 1})" if r > 4 else 0)
    for c in (et, suma):
        c.font, c.border, c.fill = HEADER_FONT, CAJA, AMARILLO
        c.alignment = Alignment(horizontal="right", vertical="center")
    suma.number_format = MONEY
    for col, ancho in ANCHOS.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "B4"
    return total


def _hoja_revisar(ws, sin_destino, sin_revisar, periodo):
    ws.cell(row=1, column=2, value="A REVISAR").font = TITLE_FONT
    ws.cell(row=2, column=2, value=f"Gastos {periodo} - comprobantes con datos a confirmar").font = BODY_FONT
    cab = ["N°", "PROVEEDOR", "COMPROBANTE", "FECHA FACT.", "IMPORTE", "DESTINO", "MEDIO", "QUÉ HAY QUE REVISAR"]
    for i, h in enumerate(cab):
        c = ws.cell(row=4, column=2 + i, value=h)
        c.font, c.border, c.fill = HEADER_FONT, CAJA, NARANJA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = 5
    filas = [(g, "Sin destino asignado: NO está sumada en ninguna hoja. Asignarle Oficina, Infraestructura, Escuela u Honorarios.")
             for g in sin_destino]
    filas += [(g, "Llegó por el portal de proveedores y todavía no se revisó en el sistema (ya está sumada en su hoja).")
              for g in sin_revisar]
    for n, (g, motivo) in enumerate(filas, 1):
        vals = [str(n), g.razon_social, f"{g.tipo_comprobante or ''} {numero_comprobante(g.punto_venta, g.numero) or ''}".strip(),
                fecha_corta(g.fecha), float(g.importe), HOJAS.get(g.categoria, "—"), medio(g.forma_pago), motivo]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + i, value=v)
            c.font, c.border = BODY_FONT, LADOS
            c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="right" if i == 4 else None)
            if i == 4:
                c.number_format = MONEY
        r += 1
    for col in range(2, 10):
        ws.cell(row=r, column=col).border = Border(top=MEDIUM)
    for col, ancho in zip("ABCDEFGHI", (1.71, 6, 38, 26, 12, 16, 20, 14, 70)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "B5"


def generar(gastos: list, anio: int, mes: int) -> bytes:
    periodo = f"{MESES[mes - 1]} {anio}"
    por_hoja = {h: [] for h in HOJAS.values()}
    sin_destino = []
    for g in sorted(gastos, key=lambda x: (x.fecha, x.id)):
        hoja = HOJAS.get(g.categoria)
        (por_hoja[hoja] if hoja else sin_destino).append(g)
    sin_revisar = [g for g in gastos if g.categoria in HOJAS and not g.revisado]

    wb = Workbook()
    wb.remove(wb.active)
    totales = {}
    for hoja, lista in por_hoja.items():
        totales[hoja] = _hoja_destino(wb.create_sheet(hoja), lista)

    ws = wb.create_sheet("TOTALES", 0)
    ws.column_dimensions["B"].width = 31.0
    ws.column_dimensions["C"].width = 24.57
    ws.cell(row=1, column=2, value="GASTOS").font = TITLE_FONT
    ws.cell(row=2, column=2, value=periodo).font = SUBTITLE_FONT
    h1, h2 = ws.cell(row=4, column=2, value="CONCEPTO"), ws.cell(row=4, column=3, value="TOTAL")
    for c in (h1, h2):
        c.font, c.fill = TOT_HEADER_FONT, VERDE
    h1.border = Border(top=THIN, bottom=THIN, left=THIN)
    h2.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    r = 5
    for hoja in HOJAS.values():
        ws.cell(row=r, column=2, value=hoja).font = TOT_ROW_FONT
        c = ws.cell(row=r, column=3, value=f"='{hoja}'!$G${totales[hoja]}")
        c.font, c.number_format = TOT_ROW_FONT, MONEY
        r += 1
    fondo = r - 1 + 4  # el espacio en blanco del modelo
    for rr in range(5, fondo + 1):
        ultimo = THIN if rr == fondo else None
        ws.cell(row=rr, column=2).border = Border(left=THIN, bottom=ultimo)
        ws.cell(row=rr, column=3).border = Border(left=THIN, right=THIN, bottom=ultimo)
    et = ws.cell(row=fondo + 1, column=2, value="TOTAL")
    et.font = TOT_HEADER_FONT
    tot = ws.cell(row=fondo + 1, column=3, value=f"=SUM(C5:C{fondo})")
    tot.font, tot.fill, tot.border, tot.number_format = TOT_HEADER_FONT, VERDE, CAJA, MONEY

    if sin_destino or sin_revisar:
        _hoja_revisar(wb.create_sheet("A REVISAR", 1), sin_destino, sin_revisar, periodo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
